import os
import io
from typing import List

from fastapi import APIRouter, UploadFile, File, HTTPException
from datetime import datetime

from langchain_core.documents import Document
from api.schema import RAGRequest, QueryResponse, SourceInfo
from processes.langchain_chain.chain import rag_chain, get_sources_info

from config.project_config import SETTINGS
from scripts.create_langchain_index import compute_doc_id
from scripts.create_langchain_index import compute_doc_id, ingest_initial_documents, ingest_new_documents



import logging
logger = logging.getLogger("router_langchain")

router = APIRouter()

collection_name = SETTINGS.qdrant_collection
qdrant_client = SETTINGS.qdrant_client


@router.get("/stats")
async def get_stats():
    """
    Endpoint para obtener estadísticas de la base de datos vectorial.
    """
    logger.info("Petición recibida: /stats")
    try:
        stats = qdrant_client.get_collection(collection_name)
        logger.info("Estadísticas obtenidas correctamente")
        return {
            "status": "ok",
            "collection_name": collection_name,
            "vectors_count": stats.vectors_count if stats.vectors_count is not None else stats.indexed_vectors_count,
            "points_count": stats.points_count,
            "segments_count": stats.segments_count,
            "optimizer_status": str(stats.optimizer_status.value) if hasattr(stats.optimizer_status, "value") else str(stats.optimizer_status),
            "collection_status": str(stats.status.value) if hasattr(stats.status, "value") else str(stats.status),
        }
    except Exception as e:
        logger.error(f"Error al obtener estadísticas: {e}")
        return {
            "status": "error",
            "message": f"Error al obtener estadísticas"
        }
        
@router.get("/health")
async def health_check():
    """
    Endpoint to check the health of the API and database connection.
    """
    try:
        # Intentamos consultar la colección en Qdrant
        try:
            qdrant_client.get_collection(collection_name)
            db_status = True
        except Exception:
            db_status = False

        return {
            "status": "ok" if db_status else "error",
            "database": "connected" if db_status else "disconnected",
            "message": "API is healthy" if db_status else "API is running but cannot reach Qdrant"
        }
    except Exception as e:
        return {
            "status": "error",
            "database": "disconnected",
            "message": f"Health check failed: {str(e)}"
        }

@router.post("/rag", response_model=QueryResponse)
async def rag_endpoint(request: RAGRequest):
    """
    Endpoint para la consulta RAG utilizando LangChain.
    """
    k = request.k_docs if request.k_docs is not None else SETTINGS.k_docs
    threshold = request.threshold if request.threshold is not None else SETTINGS.threshold

    result = await rag_chain.ainvoke({
        "question": request.question,
        "k_docs": k,
        "threshold": threshold
    })

    sources = [
        SourceInfo(source=result["source"].selection, reason=result["source"].reason)
    ] if result.get('source') else []

    return QueryResponse(
        question=result["question"],
        answer=result["answer"],
        sources=sources,
        timestamp=datetime.now()
    )

@router.post("/search")
async def search(request: RAGRequest):
    
    # si no se especifica, usar valores por defecto
    k = request.k_docs if request.k_docs is not None else SETTINGS.k_docs
    threshold = request.threshold if request.threshold is not None else SETTINGS.threshold

    # obtenemos la
    sources_info = get_sources_info(request.question, k=k, threshold=threshold)
    
    return sources_info

@router.post("/ingest/initial")
async def ingest_initial():
    """
    Fuerza la re-creación de la colección en Qdrant y reindexa
    todos los documentos iniciales (info procesada + históricos de precios).
    """
    logger.info("Petición recibida: /ingest/initial")
    try:
        total = ingest_initial_documents()
        logger.info(f"Ingesta inicial completada. Documentos indexados: {total}")
    except Exception as e:
        logger.error(f"Error durante la ingesta inicial: {e}")
        raise HTTPException(
            status_code=500,
            detail=f"Error durante la ingesta inicial: {e}",
        )

    return {
        "status": "ok",
        "message": "Ingesta inicial completada y colección recreada.",
        "indexed_documents": total,
        "collection": collection_name,
    }


ALLOWED_EXTENSIONS = {".txt", ".pdf", ".docx"}


def _extract_text_from_upload(file: UploadFile) -> str:
    """
    Extrae texto de un UploadFile en función de la extensión.
    """
    import pdfplumber
    from docx import Document as DocxDocument
    from pypdf import PdfReader

    ext = os.path.splitext(file.filename)[1].lower()

    # Leemos el contenido en memoria
    content = file.file.read()

    if ext == ".txt":
        return content.decode("utf-8", errors="ignore")

    if ext == ".pdf":
        # Opción 1: pdfplumber
        text_chunks = []
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                text_chunks.append(page.extract_text() or "")
        return "\n\n".join(text_chunks).strip()

    if ext == ".docx":
        doc = DocxDocument(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs).strip()

    # Por si acaso
    return ""


@router.post("/ingest/documents")
async def ingest_documents(files: List[UploadFile] = File(...)):
    """
    Sube uno o varios documentos (.txt, .pdf, .docx) y los añade
    a la colección existente en Qdrant sin recrearla.
    """
    logger.info("Petición recibida: /ingest/documents")
    if not files:
        logger.warning("No se ha subido ningún fichero.")
        raise HTTPException(status_code=400, detail="No se ha subido ningún fichero.")

    documents: List[Document] = []

    for f in files:
        ext = os.path.splitext(f.filename)[1].lower()
        if ext not in ALLOWED_EXTENSIONS:
            logger.warning(f"Extensión no permitida: {ext}")
            raise HTTPException(
                status_code=400,
                detail=f"Extensión no permitida: {ext}. Permitidas: {', '.join(ALLOWED_EXTENSIONS)}",
            )

        text = _extract_text_from_upload(f)
        if not text.strip():
            logger.warning(f"No se pudo extraer texto útil del archivo: {f.filename}")
            continue

        doc_id = compute_doc_id(text, f.filename)
        metadata = {
            "_collection_name": collection_name,
            "source": "info",
            "filename": f.filename,
            "uploaded_at": datetime.utcnow().isoformat(),
            "doc_id": doc_id,
        }
        documents.append(Document(page_content=text, metadata=metadata))

    if not documents:
        logger.warning("No se ha podido extraer texto de ninguno de los ficheros.")
        raise HTTPException(
            status_code=400,
            detail="No se ha podido extraer texto de ninguno de los ficheros.",
        )

    try:
        indexed = ingest_new_documents(documents)
        logger.info(f"Documentos indexados correctamente: {indexed}")
    except Exception as e:
        logger.error(f"Error durante la ingesta de documentos: {e}")
        raise HTTPException(
            status_code=500,
            detail=f"Error durante la ingesta de documentos: {e}",
        )

    return {
        "status": "ok",
        "message": "Documentos indexados correctamente.",
        "indexed_documents": indexed,
        "collection": collection_name,
    }
